"""设备管理 API：CRUD + Excel 导入导出。"""
import io

from fastapi import APIRouter, Depends, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import Device
from ..schemas import DeviceIn, DeviceOut

router = APIRouter(prefix="/api/devices", tags=["devices"])

HEADERS = ["设备名", "IP地址", "端口", "协议", "设备类型", "分组", "凭据", "SNMP模板", "描述", "启用"]

# 常用 netmiko 设备类型（前端下拉用）
DEVICE_TYPES = [
    ("huawei", "华为 VRP"),
    ("hp_comware", "H3C/华三 Comware"),
    ("zte_zxros", "中兴 ZXR10"),
    ("cisco_ios", "思科 IOS"),
    ("cisco_nxos", "思科 NX-OS"),
    ("ruijie_os", "锐捷"),
    ("linux", "Linux 服务器"),
    ("juniper_junos", "Juniper JunOS"),
    ("h3c", "H3C（h3c driver）"),
]


@router.get("/types")
def list_device_types():
    return [{"value": v, "label": l} for v, l in DEVICE_TYPES]


@router.get("/groups")
def list_groups(db: Session = Depends(get_db)):
    rows = db.query(Device.group_name).distinct().all()
    return sorted({r[0] for r in rows if r[0]})


@router.get("", response_model=list[DeviceOut])
def list_devices(group: str = "", db: Session = Depends(get_db)):
    q = db.query(Device)
    if group:
        q = q.filter(Device.group_name == group)
    return q.order_by(Device.group_name, Device.name).all()


@router.post("", response_model=DeviceOut)
def create_device(body: DeviceIn, db: Session = Depends(get_db)):
    if db.query(Device).filter(Device.name == body.name).first():
        raise HTTPException(400, f"设备名 '{body.name}' 已存在")
    device = Device(**body.model_dump())
    db.add(device)
    db.commit()
    db.refresh(device)
    return device


@router.put("/{device_id}", response_model=DeviceOut)
def update_device(device_id: int, body: DeviceIn, db: Session = Depends(get_db)):
    device = db.get(Device, device_id)
    if not device:
        raise HTTPException(404, "设备不存在")
    dup = (
        db.query(Device)
        .filter(Device.name == body.name, Device.id != device_id)
        .first()
    )
    if dup:
        raise HTTPException(400, f"设备名 '{body.name}' 已被其他设备使用")
    for k, v in body.model_dump().items():
        setattr(device, k, v)
    db.commit()
    db.refresh(device)
    return device


@router.delete("/{device_id}")
def delete_device(device_id: int, db: Session = Depends(get_db)):
    device = db.get(Device, device_id)
    if not device:
        raise HTTPException(404, "设备不存在")
    db.delete(device)
    db.commit()
    return {"ok": True}


@router.get("/template")
def download_template():
    """下载 Excel 导入模板。"""
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    ws.append(["核心交换机01", "192.168.1.1", 22, "ssh", "zte_zxros", "机房A", "网络设备凭据", "中兴只读", "示例行，导入前请删除", "是"])
    ws.append(["服务器01", "192.168.1.10", 22, "ssh", "linux", "机房A", "服务器凭据", "", "", "是"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=device_template.xlsx"},
    )


@router.post("/import")
async def import_devices(file: UploadFile, db: Session = Depends(get_db)):
    """Excel 批量导入设备。已存在的同名设备跳过。"""
    if not file.filename.endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "请上传 .xlsx 文件")
    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content))
    except Exception:
        raise HTTPException(400, "Excel 文件解析失败")
    ws = wb.active

    from ..models import Credential, SnmpProfile

    cred_by_name = {c.name: c.id for c in db.query(Credential).all()}
    snmp_by_name = {s.name: s.id for s in db.query(SnmpProfile).all()}
    existing = {d.name for d in db.query(Device).all()}

    created, skipped, errors = 0, 0, []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[0]:
            continue
        try:
            name = str(row[0]).strip()
            host = str(row[1]).strip() if row[1] else ""
            port = int(row[2]) if row[2] else 22
            protocol = str(row[3]).strip().lower() if row[3] else "ssh"
            device_type = str(row[4]).strip() if row[4] else "huawei"
            group = str(row[5]).strip() if row[5] else "默认"
            cred_name = str(row[6]).strip() if row[6] else ""
            snmp_name = str(row[7]).strip() if len(row) > 7 and row[7] else ""
            desc = str(row[8]).strip() if len(row) > 8 and row[8] else ""
            enabled = str(row[9]).strip() != "否" if len(row) > 9 and row[9] else True

            if not host:
                raise ValueError("IP地址为空")
            if name in existing:
                skipped += 1
                continue
            cred_id = cred_by_name.get(cred_name)
            if cred_name and cred_id is None:
                raise ValueError(f"凭据 '{cred_name}' 不存在，请先在凭据管理中创建")
            snmp_id = snmp_by_name.get(snmp_name)
            if snmp_name and snmp_id is None:
                raise ValueError(f"SNMP模板 '{snmp_name}' 不存在，请先创建")

            db.add(Device(
                name=name, host=host, port=port, protocol=protocol,
                device_type=device_type, group_name=group,
                credential_id=cred_id, snmp_profile_id=snmp_id,
                description=desc, enabled=enabled,
            ))
            existing.add(name)
            created += 1
        except (ValueError, TypeError) as e:
            errors.append(f"第{idx}行: {e}")
    db.commit()
    return {"created": created, "skipped": skipped, "errors": errors}


@router.get("/export")
def export_devices(db: Session = Depends(get_db)):
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for d in db.query(Device).order_by(Device.group_name, Device.name).all():
        ws.append([
            d.name, d.host, d.port, d.protocol, d.device_type, d.group_name,
            d.credential.name if d.credential else "",
            d.snmp_profile.name if d.snmp_profile else "",
            d.description,
            "是" if d.enabled else "否",
        ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=devices.xlsx"},
    )
