"""SNMP 批量表格采集：ARP/MAC/接口清单 → Excel 导出。"""
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse
from openpyxl import Workbook
from pydantic import BaseModel
from sqlalchemy.orm import Session

from ..config import DATA_DIR, get_settings
from ..core import cmdfilter, snmp
from ..core.collector import decrypt
from ..database import get_db
from ..models import Device

router = APIRouter(prefix="/api/snmp", tags=["snmp-collect"])

EXPORTS_DIR = DATA_DIR / "exports"
EXPORTS_DIR.mkdir(parents=True, exist_ok=True)

TABLE_TYPES = {
    "arp": ("ARP表", ["设备", "IP地址", "MAC地址", "接口"], snmp.collect_arp,
            lambda r: [r["_device"], r["ip"], r["mac"], r["interface"]]),
    "mac": ("MAC地址表", ["设备", "MAC地址", "VLAN", "接口"], snmp.collect_mac_table,
            lambda r: [r["_device"], r["mac"], r["vlan"], r["interface"]]),
    "interface": ("接口清单", ["设备", "接口", "别名", "状态", "速率(Mbps)", "入错包", "出错包"],
                  snmp.collect_interfaces_table,
                  lambda r: [r["_device"], r["interface"], r["alias"], r["status"],
                             r["speed_mbps"], r["err_in"], r["err_out"]]),
    "optical": ("光模块诊断", ["设备", "接口", "波长(nm)", "收光(dBm)", "发光(dBm)",
                             "温度(℃)", "电压(V)", "偏置电流(mA)"],
                snmp.collect_optical,
                lambda r: [r["_device"], r["interface"], r["wavelength_nm"],
                           r["rx_dbm"], r["tx_dbm"], r["temp_c"],
                           r["voltage_v"], r["current_ma"]]),
    "hardware": ("硬件状态", ["设备", "类型", "位置/型号", "状态", "转速", "详情"],
                 snmp.collect_hardware,
                 lambda r: [r["_device"], r["type"], r["position"], r["status"],
                            r["speed"], r["detail"]]),
}

# 内存任务表（进程重启即失效，结果文件仍在）
_jobs: dict[str, dict] = {}


class CollectIn(BaseModel):
    device_ids: list[int]
    tables: list[str]  # arp / mac / interface


def _snmp_profile_of(device: Device) -> dict | None:
    if not device.snmp_profile:
        return None
    p = device.snmp_profile
    return {
        "version": p.version, "port": p.port,
        "community": decrypt(p.community_enc),
        "v3_user": p.v3_user,
        "v3_auth_key": decrypt(p.v3_auth_key_enc),
        "v3_priv_key": decrypt(p.v3_priv_key_enc),
        "v3_auth_proto": p.v3_auth_proto,
        "v3_priv_proto": p.v3_priv_proto,
    }


def _collect_device(dev: dict, tables: list[str]) -> tuple[str, dict, str]:
    """采集单台设备的选中表格。返回 (设备名, {表类型: 行列表}, 错误)。"""
    prof = dev.get("snmp")
    if not prof:
        return dev["name"], {}, "未绑定 SNMP 模板"
    result: dict[str, list] = {t: [] for t in tables}
    for t in tables:
        _, _, func, _ = TABLE_TYPES[t]
        try:
            cmdfilter.audit("snmp", dev["name"], f"snmp-walk:{t}", True)
            rows = func(prof, dev["host"])
            for r in rows:
                r["_device"] = dev["name"]
            result[t] = rows
        except Exception as e:  # noqa: BLE001
            return dev["name"], result, f"{t}: {type(e).__name__} {e}"
    if not any(result.values()):
        return dev["name"], result, "SNMP 无数据（检查模板/网络）"
    return dev["name"], result, ""


def _run_job(job_id: str, device_ids: list[int], tables: list[str]):
    job = _jobs[job_id]
    settings = get_settings()
    from ..database import SessionLocal

    with SessionLocal() as db:
        devices = db.query(Device).filter(
            Device.id.in_(device_ids), Device.enabled.is_(True)
        ).all()
        # 会话内提取纯数据，避免 worker 线程懒加载
        dev_list = [
            {"name": d.name, "host": d.host, "snmp": _snmp_profile_of(d)}
            for d in devices
        ]

    wb = Workbook()
    wb.remove(wb.active)
    sheets = {}
    for t in tables:
        title, headers, _, _ = TABLE_TYPES[t]
        ws = wb.create_sheet(title)
        ws.append(headers)
        sheets[t] = ws

    job["total"] = len(dev_list)
    try:
        with ThreadPoolExecutor(max_workers=settings["concurrency"]) as pool:
            futures = [pool.submit(_collect_device, d, tables) for d in dev_list]
            for f in futures:
                name, rows_by_table, error = f.result()
                job["done"] += 1
                if error:
                    job["errors"].append(f"{name}: {error}")
                for t, rows in rows_by_table.items():
                    row_fn = TABLE_TYPES[t][3]
                    for r in rows:
                        sheets[t].append(row_fn(r))
    except Exception as e:  # noqa: BLE001
        job["errors"].append(f"任务异常: {type(e).__name__} {e}")

    for t in tables:
        job["counts"][t] = sheets[t].max_row - 1

    path = EXPORTS_DIR / f"snmp_{job_id}.xlsx"
    wb.save(path)
    job["file"] = path.name
    job["status"] = "done"
    job["finished_at"] = time.time()


@router.post("/collect")
def start_collect(body: CollectIn, db: Session = Depends(get_db)):
    tables = [t for t in body.tables if t in TABLE_TYPES]
    if not tables:
        raise HTTPException(400, "请至少选择一种表格（arp/mac/interface/optical/hardware）")
    if not body.device_ids:
        raise HTTPException(400, "请至少选择一台设备")
    job_id = uuid.uuid4().hex[:8]
    _jobs[job_id] = {
        "status": "running", "total": 0, "done": 0,
        "errors": [], "counts": {}, "file": None,
        "started_at": time.time(),
    }
    threading.Thread(target=_run_job, args=(job_id, body.device_ids, tables),
                     daemon=True).start()
    return {"job_id": job_id}


@router.get("/collect/{job_id}")
def collect_status(job_id: str):
    job = _jobs.get(job_id)
    if not job:
        raise HTTPException(404, "任务不存在或已过期")
    return job


@router.get("/collect/{job_id}/download")
def collect_download(job_id: str):
    job = _jobs.get(job_id)
    if not job or not job.get("file"):
        raise HTTPException(404, "文件不存在或任务未完成")
    path = EXPORTS_DIR / job["file"]
    if not path.exists():
        raise HTTPException(404, "文件已被清理")
    return FileResponse(path, filename=job["file"])
